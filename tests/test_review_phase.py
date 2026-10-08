"""The Pipeline's review phase: CV file in, parsed CV and Review out."""

from pathlib import Path

import docx
import pytest
from fakes import TAILORING, FakeModel, review_citing
from pypdf import PdfWriter

from cvxpert import config
from cvxpert.errors import ConfigurationError, ParsingError, StepError
from cvxpert.pipeline import review_phase

FIXTURES = Path(__file__).parent / "fixtures"
ALL_IDS = [f"G{n:02d}" for n in range(1, 25)]


def _pdf_cv(_tmp_path):
    return FIXTURES / "cv.pdf"  # committed: a small text PDF


def _docx_cv(tmp_path):
    document = docx.Document()
    document.add_paragraph("Jane Doe")
    document.add_paragraph("Responsible for the website.")
    path = tmp_path / "cv.docx"
    document.save(path)
    return path


def _txt_cv(tmp_path):
    path = tmp_path / "cv.txt"
    path.write_text("Jane Doe\nResponsible for the website.\n", encoding="utf-8")
    return path


@pytest.fixture
def txt_cv(tmp_path):
    return _txt_cv(tmp_path)


@pytest.mark.parametrize("make_cv", [_pdf_cv, _docx_cv, _txt_cv])
def test_pdf_docx_and_txt_cvs_each_produce_a_review(make_cv, tmp_path, use_model):
    model = use_model(FakeModel(review_citing("G09")))

    result = review_phase(make_cv(tmp_path))

    assert result.review.overall_score == 58
    assert [s.title for s in result.review.sections] == ["Experience", "Skills"]
    assert result.review.findings[0].guideline_id == "G09"
    assert "Responsible for the website." in result.cv.text
    assert "Responsible for the website." in str(model.review_requests[0])


def test_a_review_citing_every_guideline_passes_validation(txt_cv, use_model):
    # A Job Description, so G05 and G06 are kept.
    model = use_model(FakeModel(review_citing(*ALL_IDS), tailoring=TAILORING))

    result = review_phase(txt_cv, job_description="Backend Developer (Python)")

    assert [f.guideline_id for f in result.review.findings] == ALL_IDS
    assert len(model.review_requests) == 1


def test_an_unknown_guideline_id_is_retried_once_with_the_reason(txt_cv, use_model):
    model = use_model(FakeModel(review_citing("G99"), review_citing("G02")))

    result = review_phase(txt_cv)

    assert [f.guideline_id for f in result.review.findings] == ["G02"]
    assert len(model.review_requests) == 2
    assert "G99" in str(model.review_requests[1][-1])


def test_an_unknown_guideline_id_after_the_retry_is_a_clear_error(txt_cv, use_model):
    model = use_model(FakeModel(review_citing("G99")))

    with pytest.raises(StepError, match=r"review Step failed.*G99"):
        review_phase(txt_cv)
    assert len(model.review_requests) == 2


def test_an_empty_text_pdf_is_reported_as_scanned(tmp_path, use_model):
    writer = PdfWriter()
    writer.add_blank_page(width=595, height=842)
    scanned = tmp_path / "scanned.pdf"
    writer.write(scanned)
    model = use_model(FakeModel(review_citing("G09")))

    with pytest.raises(ParsingError, match="scanned PDF"):
        review_phase(scanned)
    assert model.review_requests == []


def test_an_unsupported_file_type_is_rejected(tmp_path, use_model):
    image = tmp_path / "cv.jpg"
    image.write_bytes(b"\xff\xd8")
    use_model(FakeModel(review_citing("G09")))

    with pytest.raises(ParsingError, match=r"Unsupported file type '\.jpg'"):
        review_phase(image)


def test_a_cv_over_10_mb_is_rejected(tmp_path, use_model):
    big = tmp_path / "cv.txt"
    with big.open("wb") as f:
        f.truncate(10 * 1024 * 1024 + 1)
    use_model(FakeModel(review_citing("G09")))

    with pytest.raises(ParsingError, match="larger than 10 MB"):
        review_phase(big)


def test_with_a_key_the_real_model_is_built_but_blocked_in_tests(txt_cv, monkeypatch):
    monkeypatch.setenv("OPENROUTER_API_KEY", "dummy")
    monkeypatch.setenv("CVXPERT_MODEL", "some/other-model")

    assert config.get_model().model_name == "some/other-model"
    with pytest.raises(RuntimeError, match="Model requests are not allowed"):
        review_phase(txt_cv)


def test_a_missing_api_key_fails_before_any_step_runs(tmp_path):
    # The CV does not exist: had the parse Step run first, it would report that instead.
    with pytest.raises(ConfigurationError, match="OPENROUTER_API_KEY is not set"):
        review_phase(tmp_path / "missing.txt")
